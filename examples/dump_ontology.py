# -*- coding: utf-8 -*-
"""导出 v0.0.4 动态本体全部节点与边为单个 Excel，便于人工检查质量。"""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

DB = Path(r"D:\Desktop\trial\data\ontology_v04.db")
OUT = Path(r"D:\Desktop\trial\output\ontology_v04_全部节点边.xlsx")

def prov_text(raw):
    try:
        items = json.loads(raw or "[]")
    except Exception:
        return str(raw)
    out = []
    for it in items:
        p = it.get("paper", "")
        e = (it.get("evidence") or "")[:200]
        out.append((p + " | " + e) if e else p)
    return " ;; ".join(out)

def attr_text(raw):
    try:
        return json.dumps(json.loads(raw or "{}"), ensure_ascii=False)
    except Exception:
        return str(raw)

def fmt(v):
    return v if v is not None else ""

conn = None
import sqlite3
conn = sqlite3.connect(str(DB))
conn.row_factory = sqlite3.Row

nodes = conn.execute("SELECT * FROM ontology_nodes ORDER BY node_id").fetchall()
edges = conn.execute("SELECT * FROM ontology_edges ORDER BY edge_id").fetchall()
ntypes = conn.execute("SELECT node_type, COUNT(*) c FROM ontology_nodes GROUP BY node_type ORDER BY c DESC").fetchall()
rtypes = conn.execute("SELECT relation_type, COUNT(*) c FROM ontology_edges GROUP BY relation_type ORDER BY c DESC").fetchall()
runs = conn.execute("SELECT COUNT(DISTINCT paper_key) FROM ontology_runs").fetchone()[0]

wb = Workbook()
header_fill = PatternFill("solid", fgColor="4F8CFF")
header_font = Font(color="FFFFFF", bold=True)
thin = Alignment(vertical="top", wrap_text=True)

def style(ws, widths):
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

# 概览
ws = wb.active
ws.title = "概览"
rows = [
    ["数据来源", "data/ontology_v04.db (v0.0.4 提示词，同语料50篇)"],
    ["节点总数", len(nodes)],
    ["边总数", len(edges)],
    ["已提取文献", runs],
    ["节点类型数", len(ntypes)],
    ["关系类型数", len(rtypes)],
    ["说明", "节点表含 name/aliases/attributes/confidence/来源证据；边表含主语宾语(带类型)/predicate/置信度/证据。可筛选、排序逐行检查。"],
]
for r in rows:
    ws.append(r)
ws["A1"].font = Font(bold=True); ws["B1"].font = Font(bold=True)
ws.column_dimensions["A"].width = 16
ws.column_dimensions["B"].width = 110

# 节点
ws = wb.create_sheet("节点")
ws.append(["node_id", "node_type", "name", "aliases", "confidence", "attributes", "first_seen_at", "last_seen_at", "支持论文数", "provenance(论文|证据)"])
for n in nodes:
    ws.append([n["node_id"], n["node_type"], n["name"], n["aliases"], n["confidence"],
               attr_text(n["attributes"]), n["first_seen_at"], n["last_seen_at"],
               len(json.loads(n["provenance"] or "[]")), prov_text(n["provenance"])])
style(ws, [8, 16, 42, 30, 10, 40, 19, 19, 10, 90])

# 边
ws = wb.create_sheet("边")
ws.append(["edge_id", "relation_type", "source(node)", "source_type", "target(node)", "target_type",
           "confidence", "attributes(predicate)", "provenance(论文|证据)"])
name = {n["node_id"]: n["name"] for n in nodes}
ntype = {n["node_id"]: n["node_type"] for n in nodes}
for e in edges:
    ws.append([e["edge_id"], e["relation_type"],
               name.get(e["source_node"], "?"), ntype.get(e["source_node"], "?"),
               name.get(e["target_node"], "?"), ntype.get(e["target_node"], "?"),
               e["confidence"], attr_text(e["attributes"]), prov_text(e["provenance"])])
style(ws, [8, 16, 40, 14, 40, 14, 10, 40, 80])

wb.save(OUT)
print("saved:", OUT)
print("nodes:", len(nodes), "edges:", len(edges))
conn.close()
