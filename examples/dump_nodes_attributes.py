# -*- coding: utf-8 -*-
"""导出 v0.0.4 全部节点及其属性（单文档，属性摊平）便于人工检查。"""
import json, sqlite3
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

DB = r"D:\Desktop\trial\data\ontology_v04.db"
OUT = r"D:\Desktop\trial\output\ontology_v04_节点属性检查.xlsx"

def load(x):
    try:
        return json.loads(x or "[]") if (x or "").strip().startswith(("[", "{")) else x
    except Exception:
        return x

def attr_entries(raw):
    obj = load(raw)
    if not isinstance(obj, dict):
        return [("", obj)]
    out = []
    for k, v in obj.items():
        if isinstance(v, list) and v:
            for i, item in enumerate(v):
                out.append((f"{k}[{i}]", item))
        else:
            out.append((k, v))
    return out

def val_text(v):
    if isinstance(v, dict):
        if "value" in v and "source" in v:
            return f"value={v['value']} (source: {v['source']})"
        return json.dumps(v, ensure_ascii=False)
    if isinstance(v, list):
        return json.dumps(v, ensure_ascii=False)
    return "" if v is None else str(v)

def prov_text(raw):
    items = load(raw) or []
    if not isinstance(items, list):
        return str(items)
    parts = []
    for it in items:
        if isinstance(it, dict):
            p = it.get("paper", "")
            e = (it.get("evidence") or "")[:160]
            parts.append((p + " | " + e) if e else p)
        else:
            parts.append(str(it))
    return " ;; ".join(parts)

conn = sqlite3.connect(DB)
conn.row_factory = sqlite3.Row
nodes = conn.execute("SELECT * FROM ontology_nodes ORDER BY node_id").fetchall()

wb = Workbook()
hf = PatternFill("solid", fgColor="4F8CFF"); hfont = Font(color="FFFFFF", bold=True)
wrap = Alignment(vertical="top", wrap_text=True)

def style(ws, widths):
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for cell in ws[1]:
        cell.fill = hf; cell.font = hfont
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions

ws = wb.active; ws.title = "概览"
for r in [["数据来源", "data/ontology_v04.db (v0.0.4)"], ["节点总数", len(nodes)],
          ["说明", "本表=节点主表；“属性明细”表把每个节点 attributes 逐键展开（含嵌套/来源），便于检查属性质量与配方/数值"]]:
    ws.append(r)
ws.column_dimensions["A"].width = 14; ws.column_dimensions["B"].width = 110

ws = wb.create_sheet("节点")
ws.append(["node_id", "node_type", "name", "aliases", "confidence", "first_seen_at",
           "last_seen_at", "支持论文数", "provenance(论文|证据)", "attributes(原始JSON)"])
for n in nodes:
    ws.append([n["node_id"], n["node_type"], n["name"], n["aliases"], n["confidence"],
               n["first_seen_at"], n["last_seen_at"],
               len(load(n["provenance"]) or []), prov_text(n["provenance"]),
               json.dumps(load(n["attributes"]), ensure_ascii=False)])
style(ws, [8, 16, 40, 26, 10, 18, 18, 10, 70, 80])

ws = wb.create_sheet("属性明细")
ws.append(["node_id", "node_type", "name", "属性键", "属性值(可读)", "值类型", "source(如结构内有)"])
count = 0
for n in nodes:
    for k, v in attr_entries(n["attributes"]):
        src = v.get("source") if isinstance(v, dict) else ""
        count += 1
        ws.append([n["node_id"], n["node_type"], n["name"], k, val_text(v),
                   type(v).__name__, src])
style(ws, [8, 16, 40, 26, 60, 10, 30])
ws2 = wb["概览"]
ws2.append(["属性明细行数", count])

wb.save(OUT)
print("saved:", OUT)
print("nodes:", len(nodes), "attr rows:", count)
conn.close()
