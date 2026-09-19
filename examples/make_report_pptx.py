# -*- coding: utf-8 -*-
"""生成《科研文献智能体 Agent 进度汇报》PPT。用法: bundled_python make_report_pptx.py"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BG = RGBColor(0x0F, 0x14, 0x20)
ACCENT = RGBColor(0x4F, 0x8C, 0xFF)
GREEN = RGBColor(0x2E, 0xCC, 0x8F)
ORANGE = RGBColor(0xF5, 0xA6, 0x23)
RED = RGBColor(0xFF, 0x5A, 0x6E)
TEXT = RGBColor(0xDB, 0xE4, 0xF5)
MUTED = RGBColor(0x8A, 0xA0, 0xC0)
FONT = "微软雅黑"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def set_font(run, size=16, bold=False, color=TEXT):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_slide(bg=BG):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    return s


def textbox(slide, l, t, w, h):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    return tf


def para(tf, text, size=16, bold=False, color=TEXT, level=0, bullet=False, space=6, first=False):
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.level = level
    p.space_after = Pt(space)
    r = p.add_run()
    r.text = ("• " if bullet else "") + text
    set_font(r, size=size, bold=bold, color=color)
    return p


def title_bar(slide, title, sub=None):
    tf = textbox(slide, 0.6, 0.4, 12.2, 1.1)
    para(tf, title, size=28, bold=True, color=ACCENT, first=True)
    if sub:
        para(tf, sub, size=13, color=MUTED)
    ln = slide.shapes.add_shape(1, Inches(0.6), Inches(1.35), Inches(12.1), Pt(2.2))
    ln.fill.solid(); ln.fill.fore_color.rgb = ACCENT; ln.line.fill.background()
    ln.shadow.inherit = False


def add_table(slide, rows, cols, l, t, w, h, data, header=True, sizes=None):
    nrows = len(data)
    ncols = max(len(r) for r in data)
    data = [list(r) + [''] * (ncols - len(r)) for r in data]
    shape = slide.shapes.add_table(nrows, ncols, Inches(l), Inches(t), Inches(w), Inches(h))
    tbl = shape.table
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_top = cell.margin_bottom = Pt(3)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run(); r.text = str(val)
            if header and i == 0:
                set_font(r, size=sizes[0] if sizes else 13, bold=True, color=RGBColor(0,0,0))
                cell.fill.solid(); cell.fill.fore_color.rgb = ACCENT
            else:
                set_font(r, size=sizes[1] if sizes else 12, color=TEXT)
                cell.fill.solid()
                cell.fill.fore_color.rgb = BG if i % 2 else RGBColor(0x17,0x1E,0x2E)
    return tbl


# 1 封面
s = add_slide()
tf = textbox(s, 1.0, 2.0, 11.3, 3.2)
para(tf, "科研文献智能体 Agent", size=44, bold=True, color=RGBColor(255,255,255), first=True)
para(tf, "基于 LangGraph 的三节点流水线 · 动态本体 · 提示词迭代", size=20, color=ACCENT)
para(tf, "进度汇报与问题复盘  |  版本 v0.0.1 → v0.0.4", size=15, color=MUTED)
tf2 = textbox(s, 1.0, 6.2, 11.3, 0.6)
para(tf2, "2026-09 · research-agent · github.com/Shan-Shu/trial", size=12, color=MUTED, first=True)

# 2 目录
s = add_slide()
title_bar(s, "目录")
items = [
    "一、项目目标与系统架构",
    "二、已交付功能",
    "三、提示词版本演进 v0.0.1 → v0.0.4",
    "四、量化实验与对比结果",
    "五、遇到的问题与对策",
    "六、当前成果与下一步计划",
]
tf = textbox(s, 1.2, 2.0, 10, 4.5)
for i, it in enumerate(items):
    para(tf, it, size=20, bold=(i == 0 and False), color=TEXT, first=(i == 0))
    tf.paragraphs[-1].space_after = Pt(16)

# 3 目标与架构
s = add_slide()
title_bar(s, "一、项目目标与系统架构")
tf = textbox(s, 0.6, 1.7, 12.2, 2.0)
para(tf, "目标：以多模型 LLM 驱动，从公开文献（当前 PubMed）自动完成 检索 → 质量评估 → 知识提取，构建可持续演化的科研动态本体。", size=15, first=True)
para(tf, "三个智能体节点：文献检索(DeepSeek V4) · 质量评估(GLM 4.7 Flash) · 知识提取(原 gpt-5.6，网络受限暂以 DeepSeek V4 替代)", size=14, color=MUTED)
add_table(s, 5, 3, 0.8, 3.7, 11.8, 2.6, [
    ["阶段", "职责", "输出"],
    ["检索节点", "PubMed/arXiv 检索、PDF/全文入库、元数据补全、页眉页脚清洗", "精校文献 + 监控接口"],
    ["质量节点", "权威性A×时效性T → Q=0.6A+0.4T，按 0.8/0.5 阈值路由", "A/T/Q 与路由决策"],
    ["知识节点", "分句分段 → LLM 抽取 实体/关系/属性/事件 → 写入动态本体", "可合并、可溯源的图谱"],
], sizes=(14, 12))

# 4 已交付功能
s = add_slide()
title_bar(s, "二、已交付功能")
add_table(s, 6, 2, 0.7, 1.7, 12, 5.2, [
    ["模块", "能力"],
    ["流水线 (LangGraph)", "检索→质量→知识 状态机；元数据回补循环；人工审核；断点续跑"],
    ["动态本体 (SQLite)", "节点/边去重合并、别名归一、关系同义归一、类型动态注册、溯源与版本"],
    ["图形化看板", "本体图谱(vis-network) / 智能体工作台 / 文献详情 / 6s 自动刷新"],
    ["文献源", "PubMed(NCBI+Europe PMC)、arXiv；OA 全文/摘要回退；实时监控接口"],
    ["多模型绑定", "DeepSeek V4 / GLM 4.7 Flash / (gpt-5.6→DeepSeek V4 临时替代)，.env 可切换"],
    ["工程化", "Git 版本化(v0.0.1-0.0.4)、离线测试、独立对比库、输出报表"],
], sizes=(13, 12))

# 5 版本演进
s = add_slide()
title_bar(s, "三、提示词版本演进（解决的核心问题）")
add_table(s, 5, 3, 0.7, 1.8, 12, 4.6, [
    ["版本", "提示词/工程变化", "解决什么问题"],
    ["v0.0.1", "基础三节点 + 动态本体骨架", "建立“检索→评估→提取→本体”闭环"],
    ["v0.0.2", "受控关系词表、同义归一、库内实体复用、质量评分标尺", "孤岛/断裂子图/同一关系多种写法"],
    ["v0.0.3", "实体名词化硬性规则、配方细节保留(规则8-10)、防衔接语过滤器", "“These results suggest…”等垃圾实体、过度合并风险"],
    ["v0.0.4", "事件节点名词化、trigger入attributes、participants≤5、involves克制(规则11-12)", "整句当事件导致 involves 虚高、图谱噪声"],
], sizes=(12, 13, 12))
tf = textbox(s, 0.7, 6.5, 12, 0.8)
para(tf, "远端：github.com/Shan-Shu/trial  (tags: v0.0.1~v0.0.4)", size=13, color=ACCENT, first=True)

# 6 对比指标
s = add_slide()
title_bar(s, "四、量化实验：同语料 50 篇 · 四版本对比")
add_table(s, 8, 5, 0.7, 1.8, 12.1, 4.8, [
    ["指标", "v0.0.1 旧", "v0.0.2", "v0.0.3", "v0.0.4"],
    ["节点 / 边", "1135/1164", "574/874", "569/754", "541/697"],
    ["孤立节点率", "14.5%", "0%", "1.8%", "0%"],
    ["连通分量数", "246", "9", "24", "12"],
    ["最大分量占比", "42.6%", "96.7%", "91.6%", "94.1%"],
    ["involves 占比", "50.4%", "41.9%", "31.3%", "30.1%"],
    ["伪事件(整句)节点", "154", "122", "75", "0"],
    ["关系类型(非受控)", "77(142)", "15(0)", "20(0)", "20(0)"],
], sizes=(12, 11))
tf = textbox(s, 0.7, 6.6, 12.2, 0.8)
para(tf, "另有：单主题冲量任务曾达 1060 节点；50 篇批量耗时约 40 分钟(4线程/单窗口)。", size=12, color=MUTED, first=True)

# 7 问题与对策(上)
s = add_slide()
title_bar(s, "五、遇到的问题与对策（1/2：图谱质量）")
items = [
    ("孤岛与断裂子图", "旧提示词无连通约束 → v0.0.2 起：实体复用上下文+连通性硬性要求；孤立率 14.5%→0%"),
    ("关系同义/表述不一", "无受控词表 → 受控关系词表 + 入库同义归一；关系类型 77→15→20(全部受控)"),
    ("衔接语/证据句污染", "模型把“These results suggest…”当实体 → 硬性规则8 + 代码层 is_reporting_phrase 过滤；v0.0.3 残留 6.97%→规则后拦截并归零"),
    ("事件整句化 / involves 虚高", "事件名是一整句 → v0.0.4 事件名词化+trigger 入 attributes；伪事件 75→0"),
    ("配方细节被合并(风险)", "“复用规范名”过度合并 → 规则9：带掺杂/配比等修饰必须分实体并写入 attributes"),
]
tf = textbox(s, 0.7, 1.7, 12.2, 5.6)
for i, (k, v) in enumerate(items):
    para(tf, k, size=16, bold=True, color=ACCENT, first=(i == 0))
    para(tf, v, size=13, color=TEXT)
    tf.paragraphs[-1].space_after = Pt(10)

# 8 问题与对策(下)
s = add_slide()
title_bar(s, "五、遇到的问题与对策（2/2：运行与工程）")
items = [
    ("GLM-4.7-Flash 免费档限流", "批量时频繁 429/500 → 降级为规则评分(公式不变)，GLM 可配置切回；串行限速验证通过"),
    ("DeepSeek V4-Pro 推理慢", "单篇全文 200-350s → 批量用 deepseek-v4-flash；Pro 保留可切换"),
    ("OpenAI 网络不可达", "api.openai.com 超时 → 知识提取暂用 DeepSeek V4(Pro) 替代，一行配置切回"),
    ("PubMed PDF 获取受限", "付费墙/反爬/限流 → 优先 OA 全文XML/摘要回退，并记录 fulltext_source"),
    ("SQLite 并发写锁", "多线程批量报错 → busy_timeout + 自动重试 + 断点续跑"),
    ("外部环境/会话中断", "长任务被打断 → 已做幂等续跑(按 ontology_runs 跳过已处理)"),
]
tf = textbox(s, 0.7, 1.7, 12.2, 5.6)
for i, (k, v) in enumerate(items):
    para(tf, k, size=16, bold=True, color=ORANGE, first=(i == 0))
    para(tf, v, size=13, color=TEXT)
    tf.paragraphs[-1].space_after = Pt(10)

# 9 当前成果
s = add_slide()
title_bar(s, "六、当前成果速览")
add_table(s, 6, 2, 0.7, 1.7, 12, 4.0, [
    ["维度", "现状"],
    ["代码版本", "v0.0.1~v0.0.4 已打 tag 并推送 GitHub（Shan-Shu/trial）"],
    ["本体规模(实验)", "单主题曾达 1060 节点 / 1035 边；四版本同语料对比库齐备"],
    ["质量提升", "孤立节点 0%、伪事件 0、关系全部受控、最大连通分量 ~94%"],
    ["可视化", "交互图谱 + 指标对比页（index_compare4.html）+ 本地看板 :8000"],
    ["自动化", "PubMed 监控接口、批量断点续跑、汇报生成脚本"],
], sizes=(14, 12))
tf = textbox(s, 0.7, 6.1, 12.2, 1.2)
para(tf, "下一步建议：更大语料配方保真复查、GLM 质量回填、事件粒度与 has_property 细分、看板增加“结构层/语义层”视图。", size=14, color=GREEN, first=True)

# 10 收尾
s = add_slide()
tf = textbox(s, 1.0, 2.6, 11.3, 2.5)
para(tf, "谢谢聆听", size=40, bold=True, color=RGBColor(255,255,255), first=True)
para(tf, "代码仓库：https://github.com/Shan-Shu/trial", size=16, color=ACCENT)
para(tf, "四版本对比页：output/compare4/index_compare4.html（仓库内自带）", size=13, color=MUTED)

out = r"D:\Desktop\trial\output\科研Agent_进度汇报.pptx"
prs.save(out)
print("saved:", out)

